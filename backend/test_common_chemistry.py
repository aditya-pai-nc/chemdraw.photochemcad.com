"""Offline regression checks: python -m unittest test_common_chemistry -v."""
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock, patch

import requests

import ai_curate
import common_chemistry as cas
import processor
import pubchem

ASPIRIN = {
    "rn": "50-78-2", "name": "Aspirin", "molecularFormula": "C<sub>9</sub>H<sub>8</sub>O<sub>4</sub>",
    "molecularMass": "180.16", "smile": "O=C(O)c1ccccc1OC(C)=O",
    "inchiKey": "InChIKey=BSYNRYMUTXBXSQ-UHFFFAOYSA-N",
}
SALICYLIC = {
    "rn": "69-72-7", "name": "Salicylic acid", "molecularFormula": "C7H6O3",
    "molecularMass": "138.12", "smile": "OC(=O)c1ccccc1O",
    "inchiKey": "YGSDEFSMJLZEOE-UHFFFAOYSA-N",
}
DRAWING = {
    "Compound Name": "aspirin", "ChemDraw SMILES": "CC(=O)Oc1ccccc1C(=O)O",
    "ChemDraw InChIKey": "BSYNRYMUTXBXSQ-UHFFFAOYSA-N",
    "Formula": "C9H8O4", "Molecular Weight": 180.159,
}


def response(data=None, status=200):
    return Mock(status_code=status, ok=status < 400, json=Mock(return_value=data))


def search(*rns, count=None):
    return response({"count": len(rns) if count is None else count,
                     "results": [{"rn": rn} for rn in rns]})


class CasTests(unittest.TestCase):
    def setUp(self):
        env = patch.dict(os.environ, {"CAS_API_KEY": "test-only-key", "CHEMDRAW_CAS_ENABLED": "1",
                                      "CHEMDRAW_CAS_MAX_CANDIDATES": "5"})
        env.start()
        self.addCleanup(env.stop)
        delay = patch.object(cas.time, "sleep")
        delay.start()
        self.addCleanup(delay.stop)
        network = patch.object(cas.requests, "get")
        self.get = network.start()
        self.addCleanup(network.stop)

    def test_missing_key_and_disabled_make_no_requests(self):
        for setting, expected in [({"CAS_API_KEY": ""}, "Not configured"),
                                  ({"CHEMDRAW_CAS_ENABLED": "0"}, "Disabled")]:
            with self.subTest(expected=expected), patch.dict(os.environ, setting):
                self.assertEqual(cas.CommonChemistryClient().verify(DRAWING)["CAS Verification"], expected)
        self.get.assert_not_called()

    def test_exact_structure_normalizes_html_and_key_and_caches(self):
        self.get.side_effect = [search("50-78-2"), response(ASPIRIN)]
        client = cas.CommonChemistryClient()
        result = client.verify(DRAWING)
        self.assertEqual(result["CAS Verification"], "Verified")
        self.assertEqual(result["CAS Common Chemistry Formula"], "C9H8O4")
        self.assertEqual(result["CAS Common Chemistry InChIKey"], DRAWING["ChemDraw InChIKey"])
        self.assertEqual(result["CAS Common Chemistry Molecular Weight"], 180.16)
        self.assertEqual(result["CAS Common Chemistry Link"], "https://commonchemistry.cas.org/detail?cas_rn=50-78-2")
        self.assertEqual(client.verify(DRAWING), result)
        self.assertEqual(self.get.call_count, 2)
        request = self.get.call_args_list[0].kwargs
        self.assertEqual(request["params"], {"q": DRAWING["ChemDraw InChIKey"]})
        self.assertEqual(request["headers"]["X-API-KEY"], "test-only-key")

    def test_later_exact_candidate_wins_without_borrowing_fields(self):
        sparse = {k: v for k, v in ASPIRIN.items() if k != "molecularMass"}
        self.get.side_effect = [search("69-72-7", "50-78-2"), response(SALICYLIC), response(sparse)]
        result = cas.CommonChemistryClient().verify(DRAWING)
        self.assertEqual(result["CAS Verification"], "Verified")
        self.assertEqual(result["CAS Common Chemistry RN"], "50-78-2")
        self.assertIsNone(result["CAS Common Chemistry Molecular Weight"])

    def test_opposite_stereochemistry_is_mismatch(self):
        record = {**ASPIRIN, "smile": "N[C@@H](C)C(=O)O", "canonicalSmile": "NC(C)C(=O)O"}
        self.get.side_effect = [search("50-78-2"), response(record), search(), search()]
        result = cas.CommonChemistryClient().verify({**DRAWING, "ChemDraw SMILES": "N[C@H](C)C(=O)O"})
        self.assertEqual(result["CAS Verification"], "Mismatch")

    def test_connectivity_only_smiles_cannot_verify_stereochemistry(self):
        record = {**ASPIRIN, "smile": None, "canonicalSmile": DRAWING["ChemDraw SMILES"]}
        self.get.side_effect = [search("50-78-2"), response(record), search(), search()]
        result = cas.CommonChemistryClient().verify(DRAWING)
        self.assertEqual(result["CAS Verification"], "Not comparable")

    def test_inchi_used_when_smile_missing(self):
        record = {**ASPIRIN, "smile": None,
                  "inchi": "InChI=1S/C9H8O4/c1-6(10)13-8-5-3-2-4-7(8)9(11)12/h2-5H,1H3,(H,11,12)"}
        self.get.side_effect = [search("50-78-2"), response(record)]
        self.assertEqual(cas.CommonChemistryClient().verify(DRAWING)["CAS Verification"], "Verified")

    def test_cas_number_fallback_validates_checksum_and_structure(self):
        self.get.side_effect = [search(), response(ASPIRIN)]
        result = cas.CommonChemistryClient().verify({**DRAWING, "CAS no(s)": "50-78-3; 50-78-2; 50-78-2"})
        self.assertEqual(result["CAS Verification"], "Verified")
        self.assertIn("PubChem CAS RN", result["CAS Verification Detail"])
        self.assertEqual(self.get.call_args.kwargs["params"], {"cas_rn": "50-78-2"})

    def test_name_fallback_without_pubchem(self):
        self.get.side_effect = [search(), search(), search("50-78-2"), response(ASPIRIN)]
        result = cas.CommonChemistryClient().verify(DRAWING)
        self.assertEqual(result["CAS Verification"], "Verified")
        self.assertIn("Found by Name", result["CAS Verification Detail"])

    def test_not_found_and_missing_structure_are_distinct(self):
        self.get.side_effect = [search(), search(), search()]
        self.assertEqual(cas.CommonChemistryClient().verify(DRAWING)["CAS Verification"], "Not found")
        self.get.side_effect = [search("50-78-2"), response(ASPIRIN), search()]
        self.assertEqual(cas.CommonChemistryClient().verify({**DRAWING, "ChemDraw SMILES": None})["CAS Verification"], "Not comparable")

    def test_candidate_limit_and_pagination_are_visible(self):
        with patch.dict(os.environ, {"CHEMDRAW_CAS_MAX_CANDIDATES": "1"}):
            self.get.side_effect = [search("69-72-7", "50-78-2", count=100), response(SALICYLIC)]
            result = cas.CommonChemistryClient().verify(DRAWING)
        self.assertEqual(result["CAS Verification"], "Mismatch")
        self.assertIn("additional records may exist", result["CAS Verification Detail"])
        self.assertEqual(self.get.call_count, 2)

    def test_auth_quota_server_and_network_failures_do_not_become_mismatches(self):
        for failure in [response(status=401), response(status=403), response(status=429),
                        response(status=500), requests.Timeout("sensitive request context")]:
            with self.subTest(failure=failure):
                self.get.reset_mock()
                self.get.side_effect = [failure]
                client = cas.CommonChemistryClient()
                for _ in range(2):
                    result = client.verify(DRAWING)
                    self.assertEqual(result["CAS Verification"], "Unavailable")
                    self.assertNotIn("test-only-key", str(result))
                    self.assertNotIn("sensitive", str(result))
                self.assertEqual(self.get.call_count, 1)

    def test_invalid_responses_and_incomplete_lookup_are_unavailable(self):
        for sequence in [[response([])], [response({})], [response({"results": [{"rn": 123}]})],
                         [search("50-78-2"), response({})],
                         [search("69-72-7", "50-78-2"), response(SALICYLIC), response(status=503)]]:
            with self.subTest(sequence=sequence):
                self.get.side_effect = sequence
                result = cas.CommonChemistryClient().verify(DRAWING)
                self.assertEqual(result["CAS Verification"], "Unavailable")

    def test_cas_evidence_available_for_curation_without_pubchem(self):
        row = {**DRAWING, **cas._record(ASPIRIN), "CAS Verification": "Verified"}
        evidence = ai_curate.build_evidence(row, None)
        self.assertFalse(evidence["pubchem"]["found"])
        self.assertEqual(evidence["cas_common_chemistry"]["rn"], "50-78-2")
        self.assertIn(DRAWING["ChemDraw InChIKey"], ai_curate._keys_in(evidence))

    def test_pipeline_preserves_pubchem_verdicts_curation_and_sse(self):
        def extract(_path, _mol, _images, row):
            row.update(DRAWING)
            row["Canonical SMILES"] = DRAWING["ChemDraw SMILES"]
            return {"formula": DRAWING["Formula"], "weight": DRAWING["Molecular Weight"]}

        for pubchem_hit, cas_state in [(None, "Verified"),
                                        (pubchem.PubChemHit(cid=2244, formula="C9H8O4", weight=180.16), "Unavailable")]:
            with self.subTest(cas_state=cas_state), tempfile.TemporaryDirectory() as folder:
                def enrich(_local, _name, _structures, row):
                    if pubchem_hit:
                        row["PubChem Canonical SMILES"] = DRAWING["ChemDraw SMILES"]
                    return pubchem_hit

                events = []
                cas_result = {**cas._record(ASPIRIN), "CAS Verification": cas_state, "CAS Verification Detail": "test"}
                with patch.object(processor, "_extract_structure", side_effect=extract), \
                     patch.object(processor, "_enrich_from_pubchem", side_effect=enrich), \
                     patch.object(processor, "_fill_reference_fields"), \
                     patch.object(cas.CommonChemistryClient, "verify", return_value=cas_result), \
                     patch.object(ai_curate, "is_enabled", return_value=False):
                    rows, curated = processor.process_molecules(["aspirin.cdxml"], folder + "/mol", folder + "/images", 1, events.append)
                self.assertEqual(rows[0]["CAS Verification"], cas_state)
                self.assertEqual(rows[0]["InChIKey Match?"], "✅" if pubchem_hit else "—")
                self.assertEqual(rows[0]["Match?"], "✅" if pubchem_hit else "❌")
                self.assertEqual(len(curated), 0 if pubchem_hit else 1)
                if curated:
                    self.assertEqual(curated[0]["CAS Common Chemistry RN"], "50-78-2")
                event = next(event for event in events if event["type"] == "compound")
                self.assertEqual(event["casVerification"], cas_state)
                self.assertEqual(event["casRn"], "50-78-2")

    def test_workbook_includes_cas_result_context_and_description(self):
        import pipeline
        import openpyxl

        row = {**processor._blank_row("aspirin", ""), **DRAWING, **cas._record(ASPIRIN),
               "CAS Verification": "Verified", "CAS Verification Detail": "Canonical isomeric SMILES agree."}
        curated = processor._curated_row(row, {"error": "Curation disabled"}, "No PubChem record")
        with tempfile.TemporaryDirectory() as folder:
            events = []
            with patch("chemdraw_com.open_document_count", return_value=0), \
                 patch("cdx_to_cdxml.automate_chemdraw_conversion_to_cdxml", side_effect=lambda _src, dest: Path(dest).touch()), \
                 patch("cdxml_to_ind.split_cdxml", return_value=["aspirin.cdxml"]), \
                 patch.object(processor, "process_molecules", return_value=([row], [curated])), \
                 patch.object(ai_curate, "is_enabled", return_value=False):
                pipeline.run_full_pipeline("aspirin.cdx", folder, events.append)
            workbook = openpyxl.load_workbook(Path(folder) / "aspirin_compounds.xlsx")
            self.assertEqual(workbook.sheetnames, ["Compounds", "Unmatched - AI curated", "Description"])
            for name in workbook.sheetnames[:2]:
                headers, data = list(workbook[name].values)[:2]
                record = dict(zip(headers, data))
                self.assertEqual(record["CAS Verification"], "Verified")
                self.assertEqual(record["CAS Common Chemistry RN"], "50-78-2")
            self.assertIn("CAS attribution", [r[0] for r in workbook["Description"].values])
            self.assertEqual(events[-1]["casVerifiedCount"], 1)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
