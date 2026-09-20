"""Offline CAS parser, source selection, API, export and cancellation tests."""
import asyncio
import csv
from datetime import datetime
import io
from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import Mock, patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

import bulk_cas as bulk
import bulk_cas_jobs as jobs

CAS = {"status": "Found", "rn": "50-78-2", "name": "Aspirin", "formula": "C9H8O4", "weight": 180.16,
       "smiles": "CC(=O)Oc1ccccc1C(=O)O", "inchi": "InChI=1S/C9H8O4/example",
       "inchikey": "BSYNRYMUTXBXSQ-UHFFFAOYSA-N", "url": "https://commonchemistry.cas.org/detail?cas_rn=50-78-2"}
PC = {**CAS, "cid": 2244, "url": "https://pubchem.ncbi.nlm.nih.gov/compound/2244"}


def extractor():
    client = bulk.BulkExtractor()
    client.cas.lookup_rn = Mock(return_value=CAS)
    client.pubchem.lookup = Mock(return_value=PC)
    return client


def excel_bytes(rows):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


class ParserTests(unittest.TestCase):
    def test_text_csv_tsv_unicode_and_utf16(self):
        for name, data in [
            ("input.txt", b'CAS RN\n50-78-2\n58-08-2\n'),
            ("input.txt", b'50-78-2 58-08-2\n'),
            ("input.txt", b'50-78-2,58-08-2\n'),
            ("input.csv", b'Name,CAS Number,Notes\nAspirin,50-78-2,ok\nCaffeine,58-08-2,ok\n'),
            ("input.csv", b'Name;CAS RN\nAspirin;50-78-2\nCaffeine;58-08-2\n'),
            ("input.tsv", b'Name\tCAS No.\nAspirin\t50-78-2\nCaffeine\t58-08-2\n'),
            ("input.txt", 'CAS RN\n50–78–2\n58-08-2'.encode('utf-16')),
        ]:
            with self.subTest(name=name, data=data):
                self.assertEqual([e["cas_rn"] for e in bulk.parse_upload(name, data)], ["50-78-2", "58-08-2"])

    def test_invalid_and_duplicate_entries_preserved(self):
        entries = bulk.parse_upload("input.txt", b'# list\n50-78-2\n50-78-3\nnot a CAS\n50-78-2\n')
        self.assertEqual([e["cas_rn"] for e in entries], ["50-78-2", None, None, "50-78-2"])
        self.assertEqual([e["index"] for e in entries], [1, 2, 3, 4])

    def test_hundred_entries_and_limits(self):
        self.assertEqual(len(bulk.parse_upload("list.txt", b'50-78-2\n' * 100)), 100)
        for name, data in [("list.txt", b'50-78-2\n' * 501), ("list.txt", b''),
                           ("list.pdf", b'50-78-2'), ("list.csv", b'x' * (bulk.MAX_UPLOAD_BYTES + 1)),
                           ("list.xlsx", b'not a workbook')]:
            with self.subTest(name=name, length=len(data)), self.assertRaises(ValueError):
                bulk.parse_upload(name, data)

    def test_xlsx_header_and_excel_date_error(self):
        data = excel_bytes([["Name", "CAS RN"], ["aspirin", "50-78-2"], ["mistake", datetime(2020, 5, 2)]])
        entries = bulk.parse_upload("list.xlsx", data)
        self.assertEqual(entries[0]["cas_rn"], "50-78-2")
        self.assertIsNone(entries[1]["cas_rn"])
        self.assertIn("stored this entry as a date", entries[1]["error"])

    def test_real_legacy_xls(self):
        entries = bulk.parse_upload("list.xls", (Path(__file__).parent / "tests/fixtures/cas_list.xls").read_bytes())
        self.assertEqual([e["cas_rn"] for e in entries], ["50-78-2", "58-08-2"])


class ExtractionTests(unittest.TestCase):
    def test_duplicates_cached_and_invalid_never_queried(self):
        client = extractor()
        rows = [client.extract(e) for e in bulk.parse_upload("list.txt", b'50-78-2\nwrong\n50-78-2')]
        client.cas.lookup_rn.assert_called_once_with("50-78-2")
        self.assertEqual(client.pubchem.lookup.call_count, 1)
        self.assertEqual(rows[1]["status"], "Invalid CAS")
        self.assertEqual(rows[2]["index"], 3)

    def test_source_provenance_and_missing_fields_not_mixed(self):
        client = extractor()
        client.cas.lookup_rn.return_value = {**CAS, "weight": None}
        result = client.extract(bulk.parse_upload("list.txt", b'50-78-2')[0])
        self.assertEqual(result["source"], "CAS Common Chemistry")
        self.assertIsNone(result["weight"])
        self.assertEqual(result["pubchem"]["weight"], 180.16)

    def test_pubchem_fallback_and_unavailable_distinct_from_not_found(self):
        entry = bulk.parse_upload("list.txt", b'50-78-2')[0]
        client = extractor()
        client.cas.lookup_rn.return_value = {"status": "Not found"}
        self.assertEqual(client.extract(entry)["source"], "PubChem")
        for state in ["Not found", "Unavailable"]:
            client = extractor()
            client.cas.lookup_rn.return_value = {"status": state}
            client.pubchem.lookup.return_value = {"status": "Not found"}
            self.assertEqual(client.extract(entry)["status"], state)

    def test_different_structures_flagged(self):
        client = extractor()
        client.pubchem.lookup.return_value = {**PC, "smiles": "CCO"}
        result = client.extract(bulk.parse_upload("list.txt", b'50-78-2')[0])
        self.assertEqual(result["agreement"], "Differ")
        self.assertIn("Review both source records", result["notes"])

    def test_pubchem_selects_matching_candidate_without_merging(self):
        props = [{"CID": 702, "SMILES": "CCO", "MolecularWeight": "46.07"},
                 {"CID": 2244, "SMILES": CAS["smiles"], "InChI": CAS["inchi"]}]
        response = Mock(status_code=200, ok=True, json=Mock(return_value={"PropertyTable": {"Properties": props}}))
        with patch.object(bulk.requests, "get", return_value=response), patch.object(bulk.pubchem, "_rate_limit"):
            result = bulk.PubChemCasClient().lookup("50-78-2", CAS)
        self.assertEqual(result["cid"], 2244)
        self.assertIsNone(result["weight"])
        self.assertEqual(result["candidate_cids"], [702, 2244])

    def test_malformed_pubchem_record_is_unavailable(self):
        response = Mock(status_code=200, ok=True, json=Mock(return_value={"PropertyTable": {"Properties": [None]}}))
        with patch.object(bulk.requests, "get", return_value=response), patch.object(bulk.pubchem, "_rate_limit"):
            result = bulk.PubChemCasClient().lookup("50-78-2", CAS)
        self.assertEqual(result["status"], "Unavailable")

    def test_export_fields_source_sheets_and_literal_excel_text(self):
        client = extractor()
        entries = bulk.parse_upload("list.txt", b'50-78-2\n=HYPERLINK("bad")')
        rows = [client.extract(entry) for entry in entries]
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            bulk.write_exports(rows, folder, 2, "done")
            workbook = load_workbook(folder / "bulk_cas_results.xlsx")
            self.assertEqual(workbook.sheetnames, ["Compounds", "Source records", "Read me"])
            header, first, second = list(workbook["Compounds"].values)
            result = dict(zip(header, first))
            for field in ["Molecular Formula", "Molecular Weight", "SMILES", "InChI", "PubChem Link", "CAS Common Chemistry Link"]:
                self.assertTrue(result[field])
            self.assertEqual(second[1], '=HYPERLINK("bad")')
            self.assertEqual(workbook["Compounds"]["B3"].data_type, "s")
            self.assertEqual(workbook["Source records"].max_row, 5)
            workbook.close()
            with (folder / "bulk_cas_results.csv").open(encoding="utf-8-sig", newline="") as fh:
                data = list(csv.DictReader(fh))
            self.assertTrue(data[1]["Input CAS"].startswith("'="))


class ApiTests(unittest.TestCase):
    def test_hundred_row_upload_progress_download_and_restart(self):
        with tempfile.TemporaryDirectory() as tmp, patch.object(jobs, "BulkExtractor", side_effect=extractor):
            manager = jobs.BulkManager(Path(tmp))
            app = FastAPI()
            app.include_router(jobs.router)
            with patch.object(jobs, "manager", manager), TestClient(app) as client:
                response = client.post("/api/bulk-cas", files={"file": ("list.csv", b'CAS RN\n' + b'50-78-2\n' * 100)})
                self.assertEqual(response.status_code, 200)
                job_id = response.json()["job_id"]
                deadline = time.monotonic() + 10
                while time.monotonic() < deadline:
                    snapshot = client.get(f"/api/bulk-cas/{job_id}").json()
                    if snapshot["status"] in jobs.TERMINAL:
                        break
                    time.sleep(0.02)
                self.assertEqual(snapshot["status"], "done")
                self.assertEqual(snapshot["processed"], 100)
                self.assertEqual(snapshot["unique"], 1)
                self.assertTrue(snapshot["has_exports"])
                for format in ["xlsx", "csv"]:
                    self.assertEqual(client.get(f"/api/bulk-cas/{job_id}/download/{format}").status_code, 200)
                reloaded = jobs.BulkManager(Path(tmp)).snapshot(job_id)
                self.assertEqual(len(reloaded["rows"]), 100)
                self.assertTrue(reloaded["has_exports"])
                self.assertEqual(client.get("/api/bulk-cas/not-a-job").status_code, 404)
                self.assertEqual(client.post("/api/bulk-cas", files={"file": ("bad.pdf", b'50-78-2')}).status_code, 400)


class CancellationTests(unittest.IsolatedAsyncioTestCase):
    async def test_cancel_running_stops_after_inflight_and_exports_partial_results(self):
        with tempfile.TemporaryDirectory() as tmp:
            manager = jobs.BulkManager(Path(tmp))
            started, release = threading.Event(), threading.Event()
            real = extractor()
            def slow(entry):
                started.set()
                release.wait(5)
                return real.extract(entry)
            client = Mock(extract=Mock(side_effect=slow))
            with patch.object(jobs, "BulkExtractor", return_value=client):
                job = await manager.create("list.txt", b'50-78-2\n58-08-2')
                self.assertTrue(await asyncio.to_thread(started.wait, 2))
                await manager.cancel(job.id)
                release.set()
                await asyncio.gather(*manager.tasks)
            self.assertEqual(job.status, "cancelled")
            self.assertEqual(len(job.rows), 1)
            self.assertTrue(job.exports_ready)
            self.assertEqual(client.extract.call_count, 1)

    async def test_cancel_queued_makes_no_network_calls(self):
        with tempfile.TemporaryDirectory() as tmp:
            manager = jobs.BulkManager(Path(tmp))
            await manager.semaphore.acquire()
            with patch.object(jobs, "BulkExtractor") as client:
                job = await manager.create("list.txt", b'50-78-2')
                await manager.cancel(job.id)
                self.assertEqual(job.status, "cancelled")
                manager.semaphore.release()
                await asyncio.gather(*manager.tasks)
                client.assert_not_called()


if __name__ == "__main__":
    unittest.main()
