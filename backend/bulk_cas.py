"""CAS list parsing, source-preserving lookups, and spreadsheet exports."""
from __future__ import annotations

import csv
from datetime import date, datetime
import io
import math
import re
from pathlib import Path
from typing import Any

import requests

import common_chemistry
import pubchem
from inchi_tools import canonical_smiles

MAX_ENTRIES = 500
MAX_UPLOAD_BYTES = 2 * 1024 * 1024
EXTENSIONS = {".txt", ".csv", ".tsv", ".xlsx", ".xls"}
HEADERS = {"cas", "casrn", "casno", "casnos", "casnumber", "casnumbers", "casregistrynumber", "casregistrynumbers"}
FIELDS = ("name", "formula", "weight", "smiles", "inchi", "inchikey")


def normalize_cas(value: Any) -> str | None:
    if not isinstance(value, str):
        return None
    value = value.strip().translate(str.maketrans({"–": "-", "—": "-", "−": "-", "‑": "-"}))
    value = re.sub(r"^CAS\s*(?:RN|No\.?|Number)?\s*[:#]?\s*", "", value, flags=re.I)
    value = re.sub(r"\s*-\s*", "-", value)
    if not re.fullmatch(r"[1-9]\d{1,6}-\d{2}-\d", value):
        return None
    return value if common_chemistry._cas_rns(value) == [value] else None


def _header(value: Any) -> bool:
    return isinstance(value, str) and re.sub(r"[^a-z]", "", value.lower()) in HEADERS


def _tables(filename: str, data: bytes):
    ext = Path(filename).suffix.lower()
    if ext == ".xlsx":
        from openpyxl import load_workbook
        with io.BytesIO(data) as stream:
            workbook = load_workbook(stream, read_only=True, data_only=True)
            try:
                for sheet in workbook:
                    if sheet.max_row and sheet.max_row > 5000 or sheet.max_column and sheet.max_column > 100:
                        raise ValueError("Excel sheets must have at most 5,000 rows and 100 columns. Upload only the CAS list.")
                    yield sheet.title, list(sheet.iter_rows(values_only=True))
            finally:
                workbook.close()
    elif ext == ".xls":
        import xlrd
        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
        try:
            for sheet in workbook.sheets():
                if sheet.nrows > 5000 or sheet.ncols > 100:
                    raise ValueError("Excel sheets must have at most 5,000 rows and 100 columns. Upload only the CAS list.")
                rows = []
                for i in range(sheet.nrows):
                    rows.append([xlrd.xldate_as_datetime(c.value, workbook.datemode)
                                 if c.ctype == xlrd.XL_CELL_DATE else c.value for c in sheet.row(i)])
                yield sheet.name, rows
        finally:
            workbook.release_resources()
    else:
        encoding = "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            text = data.decode("cp1252")
        if "\x00" in text:
            raise ValueError("The upload is not a readable text file. Save as UTF-8 TXT or CSV.")
        if ext == ".tsv":
            delimiter = "\t"
        else:
            try:
                delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t").delimiter
            except csv.Error:
                delimiter = "," if ext == ".csv" else "\t"
        yield "File", list(csv.reader(io.StringIO(text), delimiter=delimiter))


def parse_upload(filename: str, data: bytes) -> list[dict]:
    if not data:
        raise ValueError("The uploaded file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("Upload a file no larger than 2 MB.")
    if Path(filename).suffix.lower() not in EXTENSIONS:
        raise ValueError("Upload a TXT, CSV, TSV, XLSX or XLS file.")
    entries = []
    try:
        for sheet, rows in _tables(filename, data):
            header = next(((i, j) for i, row in enumerate(rows[:20])
                           for j, cell in enumerate(row) if _header(cell)), None)
            start, column = (header[0] + 1, header[1]) if header else (0, None)
            for row_index, row in enumerate(rows[start:], start + 1):
                cells = [row[column] if column < len(row) else None] if column is not None else row
                for cell in cells:
                    if cell is None or not str(cell).strip():
                        continue
                    text = str(cell).strip()
                    if text.startswith("#") or _header(text):
                        continue
                    tokens = re.split(r"[;,\r\n]+", text)
                    if len(tokens) == 1 and len(text.split()) > 1 and all(normalize_cas(t) for t in text.split()):
                        tokens = text.split()
                    for token in tokens:
                        token = token.strip()
                        if not token:
                            continue
                        if len(token) > 300:
                            raise ValueError("An entry is too long. Upload a CAS list or label the CAS column 'CAS RN'.")
                        rn = normalize_cas(token)
                        entries.append({"index": len(entries) + 1, "input": token, "cas_rn": rn,
                                        "location": f"{sheet}, row {row_index}",
                                        "error": None if rn else (
                                            "Excel stored this entry as a date. Format the CAS column as Text and re-enter the number."
                                            if isinstance(cell, (date, datetime)) else
                                            "Invalid CAS format or check digit. Expected a number such as 50-78-2.")})
                        if len(entries) > MAX_ENTRIES:
                            raise ValueError(f"Upload at most {MAX_ENTRIES} CAS entries per file.")
    except ValueError:
        raise
    except ImportError as exc:
        raise ValueError("The Excel reader is missing. Install the backend requirements.") from exc
    except Exception as exc:
        raise ValueError("Could not read this file. Save the CAS list as TXT, CSV or a valid Excel workbook.") from exc
    if not entries:
        raise ValueError("No CAS entries found. Use one number per line or a column labelled 'CAS RN'.")
    return entries


class PubChemCasClient:
    def __init__(self):
        self.unavailable = None

    def lookup(self, rn: str, cas_record: dict) -> dict:
        result = {"status": "Not found", "note": None, "cid": None, "url": None,
                  "candidate_cids": [], **dict.fromkeys(FIELDS)}
        if self.unavailable:
            return {**result, "status": "Unavailable", "note": self.unavailable}
        try:
            pubchem._rate_limit()
            response = requests.get(f"{pubchem.BASE}/compound/name/{rn}/property/{pubchem.LOOKUP_PROPS}/JSON",
                                    timeout=(5, 20))
            if response.status_code == 404:
                return {**result, "note": "No PubChem record found for this CAS number."}
            if not response.ok:
                self.unavailable = f"PubChem is unavailable (HTTP {response.status_code})."
                return {**result, "status": "Unavailable", "note": self.unavailable}
            props = response.json()["PropertyTable"]["Properties"]
            if not isinstance(props, list) or not props or any(not isinstance(p, dict) for p in props):
                raise ValueError("Invalid PubChem properties")
            candidates = [pubchem.PubChemHit.from_properties(p, "CAS RN", rn) for p in props[:25]]
            if any(not isinstance(hit.cid, int) or hit.cid <= 0 for hit in candidates):
                raise ValueError("Invalid PubChem CID")
            # Prefer the candidate that agrees structurally with CAS. Each
            # returned field still belongs to that single PubChem record.
            cas_smiles = canonical_smiles(cas_record.get("smiles"))
            chosen = next((hit for hit in candidates if cas_smiles and canonical_smiles(hit.smiles) == cas_smiles), candidates[0])
            if chosen.weight is not None and not math.isfinite(chosen.weight):
                chosen.weight = None
            cids = [hit.cid for hit in candidates]
            note = None
            if len(props) > 1:
                note = f"{len(props)} PubChem candidates; selected CID {chosen.cid}."
            if len(props) > 25:
                note += " Only the first 25 candidates were inspected."
            return {"status": "Found", "name": chosen.title or chosen.iupac_name,
                    "formula": chosen.formula, "weight": chosen.weight, "smiles": chosen.smiles,
                    "inchi": chosen.inchi, "inchikey": chosen.inchikey,
                    "cid": chosen.cid, "candidate_cids": cids,
                    "url": f"https://pubchem.ncbi.nlm.nih.gov/compound/{chosen.cid}", "note": note}
        except (requests.RequestException, ValueError, KeyError, TypeError):
            self.unavailable = "PubChem could not be reached or returned an invalid response."
            return {**result, "status": "Unavailable", "note": self.unavailable}


class BulkExtractor:
    def __init__(self):
        self.cas = common_chemistry.CommonChemistryClient()
        self.pubchem = PubChemCasClient()
        self.cache: dict[str, dict] = {}

    def extract(self, entry: dict) -> dict:
        rn = entry["cas_rn"]
        if not rn:
            return {**entry, "status": "Invalid CAS", "notes": entry["error"], "source": None,
                    "cas": {"status": "Skipped"}, "pubchem": {"status": "Skipped"},
                    "agreement": "Not comparable", **dict.fromkeys(FIELDS)}
        if rn not in self.cache:
            cas = self.cas.lookup_rn(rn)
            pc = self.pubchem.lookup(rn, cas)
            cas_found, pc_found = cas["status"] == "Found", pc["status"] == "Found"
            primary = cas if cas_found else pc if pc_found else {}
            source = "CAS Common Chemistry" if cas_found else "PubChem" if pc_found else None
            a, b = canonical_smiles(cas.get("smiles")), canonical_smiles(pc.get("smiles"))
            agreement = "Agree" if a and b and a == b else "Differ" if a and b else "Not comparable"
            status = "Found" if source else "Unavailable" if any(s["status"] in ("Unavailable", "Not configured", "Disabled") for s in [cas, pc]) else "Not found"
            notes = [s.get("note") for s in [cas, pc] if s.get("note")]
            if agreement == "Differ":
                notes.append("CAS and PubChem structures differ. Review both source records before using the result.")
            missing = [name for name in ["formula", "weight", "smiles", "inchi"] if primary and primary.get(name) is None]
            if missing:
                notes.append(f"{source} did not supply: {', '.join(missing)}. See the other source record if available.")
            self.cache[rn] = {"status": status, "source": source, "agreement": agreement,
                              "cas": cas, "pubchem": pc, "notes": " ".join(notes) or None,
                              **{key: primary.get(key) for key in FIELDS}}
        return {**entry, **self.cache[rn]}


def flat_row(row: dict) -> dict:
    return {"Input row": row["index"], "Input CAS": row["input"], "CAS RN": row["cas_rn"],
            "Status": row["status"], "Name": row.get("name"), "Molecular Formula": row.get("formula"),
            "Molecular Weight": row.get("weight"), "SMILES": row.get("smiles"), "InChI": row.get("inchi"),
            "InChIKey": row.get("inchikey"), "Data Source": row.get("source"),
            "PubChem CID": row["pubchem"].get("cid"), "PubChem Link": row["pubchem"].get("url"),
            "CAS Common Chemistry Link": row["cas"].get("url"), "CAS Status": row["cas"]["status"],
            "PubChem Status": row["pubchem"]["status"], "Structure Agreement": row["agreement"],
            "Input Location": row["location"], "Notes": row.get("notes")}


def write_exports(rows: list[dict], folder: Path, total: int, status: str):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Compounds"
    flattened = [flat_row(row) for row in rows]
    headers = list(flat_row({"index": 0, "input": "", "cas_rn": None, "status": "", "location": "",
                             "cas": {"status": ""}, "pubchem": {"status": ""}, "agreement": ""}))

    def append(ws, values):
        ws.append(values)
        # Source names and uploaded text must remain literal spreadsheet text.
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"

    append(summary, headers)
    for row in flattened:
        append(summary, list(row.values()))
    sources = workbook.create_sheet("Source records")
    append(sources, ["Input row", "Input CAS", "Source", "Status", "Registered CAS RN", "PubChem CID",
                     "Name", "Molecular Formula", "Molecular Weight", "SMILES", "InChI", "InChIKey", "Link", "Notes"])
    for row in rows:
        for name, source in [("CAS Common Chemistry", row["cas"]), ("PubChem", row["pubchem"])]:
            append(sources, [row["index"], row["input"], name, source["status"], source.get("rn"), source.get("cid"),
                             *[source.get(key) for key in FIELDS], source.get("url"), source.get("note")])
    for ws in [summary, sources]:
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1E293B")
            ws.column_dimensions[cell.column_letter].width = 24
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top")
    readme = workbook.create_sheet("Read me")
    for values in [
        ("Bulk chemical information extractor", f"{len(rows)} of {total} input entries processed; job status: {status}."),
        ("Sources", "CAS Common Chemistry and PubChem; source values are retained independently on Source records."),
        ("Primary data", "CAS Common Chemistry is used when it returns a record; otherwise PubChem. Missing fields are left blank, never borrowed across sources or candidates."),
        ("Structure Agreement", "Canonical isomeric SMILES agree/differ, or are not comparable. Matching formulas or CAS search terms alone do not verify a structure."),
        ("PubChem candidates", "Up to 25 candidates are inspected. Prefer the one whose canonical isomeric SMILES agrees with CAS; otherwise show the first candidate and its provenance."),
        ("Input", "Input order and duplicates are preserved. Invalid CAS formats/check digits are listed without network lookups. Blank cells and comments are skipped."),
        ("CAS attribution", "CAS Common Chemistry, CAS, a division of the American Chemical Society. https://commonchemistry.cas.org/ . CC BY-NC 4.0: https://creativecommons.org/licenses/by-nc/4.0/"),
        ("CAS transformations", "HTML formatting is removed from names/formulas. SMILES may be reconstructed from InChI. Each record has its source link."),
        ("PubChem reference", "https://pubchem.ncbi.nlm.nih.gov/ ; each retrieved compound is linked separately."),
        ("Unavailable", "Network/access failures are distinct from missing records. Some substances do not have a discrete structure, SMILES or InChI."),
    ]:
        append(readme, list(values))
    readme.column_dimensions["A"].width = 30
    readme.column_dimensions["B"].width = 110
    for row in readme:
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(folder / "bulk_cas_results.xlsx")
    with (folder / "bulk_cas_results.csv").open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(headers)
        for row in flattened:
            writer.writerow(["'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")) else value
                             for value in row.values()])
